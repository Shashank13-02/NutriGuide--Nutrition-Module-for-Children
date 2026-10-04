import openpyxl
from pathlib import Path

p = Path("Text slm datasets")

for f in sorted(p.glob("*.xlsx")):
    try:
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        print(f"SUCCESS: {f.name} -> {len(wb.sheetnames)} sheets: {wb.sheetnames}")
        wb.close()
    except Exception as e:
        print(f"SKIPPED (locked/err): {f.name} -> {e}")
